# -*- coding: utf-8 -*-
"""Dump DOCX paragraphs (with index) to a text file for structure inspection."""
import sys
from docx import Document

src = sys.argv[1]
out = sys.argv[2]
doc = Document(src)
with open(out, "w", encoding="utf-8") as f:
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if t:
            f.write(f"[{i}] {t}\n")
    # tables too, in case questions live in tables
    for ti, table in enumerate(doc.tables):
        f.write(f"\n=== TABLE {ti} ===\n")
        for row in table.rows:
            f.write(" | ".join(c.text.strip() for c in row.cells) + "\n")
print("done")
